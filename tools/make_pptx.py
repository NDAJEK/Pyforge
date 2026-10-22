#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 PPT 实训素材与参考答案（python-pptx + 少量 XML 后处理）"""
import os, re, shutil, zipfile
from pptx import Presentation
from pptx.util import Emu, Pt

SLIDES = [
    ("2026 年度产品规划", "技术部 · 2025 年 12 月", None),
    ("一、市场回顾", None, ["去年整体市场增长 12%", "竞品在中端价位发力", "我们的复购率保持领先"]),
    ("二、产品路线", None, ["Q1 完成核心重构", "Q2 上线移动端", "Q3 打通数据平台"]),
    ("三、团队分工", None, ["研发：架构与实现", "设计：交互与视觉", "运营：内容与增长"]),
]

FONT = "微软雅黑"


def build(path, finished):
    prs = Presentation()
    if finished:
        prs.slide_width = Emu(12192000)   # 标准 16:9
        prs.slide_height = Emu(6858000)

    for i, (title, sub, bullets) in enumerate(SLIDES):
        layout = prs.slide_layouts[0] if i == 0 else prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        for p in slide.shapes.title.text_frame.paragraphs:
            for r in p.runs:
                r.font.name = FONT if finished else "宋体"
                r.font.size = Pt(44 if finished else 32)
        if sub and len(slide.placeholders) > 1:
            ph = slide.placeholders[1]
            ph.text = sub
            for p in ph.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = FONT if finished else "宋体"
                    r.font.size = Pt(20 if finished else 18)
        if bullets and len(slide.placeholders) > 1:
            ph = slide.placeholders[1]
            tf = ph.text_frame
            tf.text = bullets[0]
            for b in bullets[1:]:
                tf.add_paragraph().text = b
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.name = FONT if finished else "宋体"
                    r.font.size = Pt(24 if finished else 18)

    if finished:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "谢谢"
        for p in slide.shapes.title.text_frame.paragraphs:
            for r in p.runs:
                r.font.name = FONT
                r.font.size = Pt(44)

    tmp = path + ".tmp"
    prs.save(tmp)

    if not finished:
        shutil.move(tmp, path)
        print("生成", path, os.path.getsize(path), "字节")
        return

    src = zipfile.ZipFile(tmp)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if re.match(r'ppt/slides/slide\d+\.xml$', item.filename):
                x = data.decode('utf-8')
                x = x.replace('<a:latin typeface="%s"/>' % FONT,
                              '<a:latin typeface="%s"/><a:ea typeface="%s"/>' % (FONT, FONT))
                if '</p:sld>' in x and '<p:transition' not in x:
                    x = x.replace('</p:sld>', '<p:transition><p:fade/></p:transition></p:sld>')
                data = x.encode('utf-8')
            out.writestr(item, data)
    src.close()
    os.remove(tmp)
    print("生成", path, os.path.getsize(path), "字节")


if __name__ == '__main__':
    base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    d = os.path.join(base, 'office', 'task-03')
    os.makedirs(d, exist_ok=True)
    build(os.path.join(d, '产品规划-素材.pptx'), False)
    build(os.path.join(d, '产品规划-参考答案.pptx'), True)
